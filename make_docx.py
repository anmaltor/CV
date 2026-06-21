#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN = RGBColor(0x46, 0x60, 0x4f)
DARK = RGBColor(0x2b, 0x2b, 0x2b)

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.6); s.bottom_margin = Inches(0.6)
    s.left_margin = Inches(0.7); s.right_margin = Inches(0.7)

normal = doc.styles['Normal']
normal.font.name = 'Georgia'; normal.font.size = Pt(10); normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(3); normal.paragraph_format.line_spacing = 1.05

def name(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.font.size = Pt(26); r.font.color.rgb = DARK
    p.paragraph_format.space_after = Pt(0)

def headline(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = GREEN
    p.paragraph_format.space_after = Pt(8)

def section(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = GREEN
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6'); bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), '5f7d6e')
    pbdr.append(bottom); pPr.append(pbdr)

def jobtitle(title, meta):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title); r.bold = True; r.font.size = Pt(10.5)
    p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(meta); r2.bold = True; r2.font.size = Pt(9.5); r2.font.color.rgb = GREEN

def desc(text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9.5)

def subhead(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text); r.bold = True; r.italic = True; r.font.size = Pt(9.5)

def bullet(segments):
    # segments: list of (text, bold)
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    for text, bold in segments:
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(10)

def para(text, justify=True):
    p = doc.add_paragraph()
    if justify: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); r.font.size = Pt(10)

name("Antonio Mallol Torralbo")
headline("Bid & Proposal Director  —  Rail & Transit PPP / O&M")

# Contact line
c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(6)
rc = c.add_run("Toronto, Canada (open to relocation — UAE / Gulf)  •  +1 647-539-2191  •  Antonio.MallolTorralbo@gmail.com  •  linkedin.com/in/antonio-mallol-torralbo")
rc.font.size = Pt(9)

section("Profile")
para("Bid and proposal leader with 20+ years in rail and transit, holding full profit-and-loss "
     "accountability for major PPP/P3 pursuits across North America, EMEA, and Asia Pacific. Built and "
     "led Bechtel's global O&M bid function, shaping win strategy, Concept of Operations, lifecycle cost "
     "models, and risk-transfer pricing on multi-billion-dollar metro, LRT, and infrastructure tenders — "
     "repeatedly delivering proposals 20–40% below initial budget. Combines this commercial bid leadership "
     "with hands-on delivery as an operator and maintainer, giving submissions credibility that wins: every "
     "assumption is grounded in what is genuinely deliverable. Strong fit for leading the O&M bid for the "
     "Dubai Blue Line, bringing direct metro-PPP operating experience (Madrid) to RATP's pursuit.")

section("Core Strengths")
para("Bid & Proposal Management  •  O&M Tender Development  •  PPP / P3 Commercial Structuring  •  "
     "Win Strategy & Bid Governance  •  Lifecycle Cost & 30-yr P&L Modeling  •  Risk Transfer & Pricing Strategy  •  "
     "Concept of Operations (ConOps)  •  Systems Integration  •  Client & Stakeholder Engagement  •  "
     "Operational Deliverability Assurance", justify=False)

section("Experience")

jobtitle("Head of Operations & Maintenance (O&M) — Global Bids",
         "Bechtel Infrastructure  |  Reston, US & Toronto, Canada  |  Jul 2015 – Mar 2019")
desc("Established and led Bechtel's global O&M business line, driving bid pursuit and win strategy for the company's worldwide PPP and Delivery-Partner portfolio.")
bullet([("Owned all O&M proposals and rail-system bids", True), (" with full profit-and-loss accountability through the Design-Build phase, governing pricing, scope, and risk position.", False)])
bullet([("Designed the Concept of Operations (ConOps)", True), (" for major global rail bids, defining staffing models for station operations, drivers, and control-room personnel as the basis of cost and price.", False)])
bullet([("Built 30-year lifecycle cost models", True), (", quantifying the P&L impact of headway, dwell time, and fleet-utilization assumptions to optimize the operating-cost case.", False)])
bullet([("Led bid risk analysis", True), (" on new pursuits and developed Bechtel's O&M business plan and long-term strategy across North America, EMEA, and Asia Pacific.", False)])
bullet([("Recruited and trained the multidisciplinary bid team — including operational rail-modeling experts — to support PPP and Delivery-Partner pursuits.", False)])
subhead("Selected bid outcomes")
bullet([("Rail systems integration & T&C proposal (DBJV, Ottawa LRT PPP) — ", False), ("30% below", True), (" initial budget.", False)])
bullet([("O&M & rail-systems bid (OMJV, Finch LRT PPP) — ", False), ("20% below", True), (" initial budget.", False)])
bullet([("Bridge tolling, intelligent traffic control & health-monitoring bid (Gordie Howe Bridge PPP) — ", False), ("40% below", True), (" initial budget, with O&M also under estimate.", False)])
bullet([("Edmonton Valley Line LRT (P3, 30-yr O&M) — bid through to O&M company mobilization as a member of the execution Management Committee.", False)])

jobtitle("Senior Consultant — Head of Fixed Assets",
         "Boxfish Infrastructure Group Inc.  |  Toronto, Canada  |  Jul 2019 – Jan 2024")
bullet([("Delivered a maintenance bid 25% below estimate", True), (" by re-engineering the maintenance risk-transfer model for the concession.", False)])
bullet([("Developed the asset-management plan for a ", False), ("$20B network upgrade", True), (" (OnCorr), ensuring assets met regulatory codes ahead of operational handover.", False)])
bullet([("Led vehicle–infrastructure integration on the multi-billion-dollar Eglinton Crosstown, mediating technical clashes between the fleet manufacturer and the construction JV.", False)])
bullet([("Worked with the incoming operator (Deutsche Bahn) to validate ConOps and timetable, engineering Union Station capacity upgrades (24m high-speed switches, double berthing) to clear the network's primary bottleneck.", False)])

jobtitle("General Manager — Eglinton Crosstown LRT (CTSM)",
         "ACS Infrastructure  |  Toronto, Canada  |  Jan 2024 – Present")
desc("Delivering, in live operations, the O&M model that PPP bids promise — validating bid assumptions against real SLA performance.")
bullet([("Accountable for ", False), ("Service Availability against ~98% SLA targets", True), (", managing the fleet/infrastructure performance interface that underpins any O&M bid case.", False)])
bullet([("Led the Trial Running & Operational Readiness strategy and the Shadow Operations handover of 25 stations and 19km of track from construction to live status.", False)])
bullet([("Developed failure-recovery SOPs and chaired the Safety Review Board, leading root-cause investigations to preserve the Safety Case and regulatory compliance.", False)])
bullet([("Manage the contractual relationship with Metrolinx, holding the line on SLAs and operational KPIs.", False)])

jobtitle("Maintenance Director — Metro Ligero Oeste (Madrid Metro PPP)",
         "Metro Ligero Oeste S.A.  |  Pozuelo de Alarcón, Spain  |  Sep 2006 – Jul 2015")
desc("Led the maintenance and rehabilitation phase of a light-rail PPP concession within the Madrid metro network.")
bullet([("Cut the initial infrastructure maintenance budget 35%", True), (" via a self-performing strategy, generating significant concession profit.", False)])
bullet([("Delivered ", False), ("99.5% Service Availability", True), (" across ML2/ML3, managing daily fleet output to peak-hour headways.", False)])
bullet([("Reduced service-affecting incidents 30% through a predictive-intervention strategy and the energy bill 30% via tunnel/street lighting and ventilation control.", False)])

section("Education")
bullet([("Executive MBA", True), (" — IESE Business School, Spain (2013); ranked top-3 globally by the Financial Times.", False)])
bullet([("Executive MBA, Lean Manufacturing", True), (" — FH Ludwigshafen, Germany (2006).", False)])
bullet([("MEng, Industrial Engineering (Mechanics)", True), (" — Polytechnic University of Valencia, Spain (2002).", False)])

section("Languages")
para("Spanish — Native    •    English — Fluent", justify=False)

doc.save("CV Antonio Mallol_Bid Manager_RATP Dubai Blue Line.docx")
print("DOCX saved")
